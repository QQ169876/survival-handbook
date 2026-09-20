# -*- coding: utf-8 -*-
"""docx 排版公共工具：中文标题、正文、项目符号、表格、插图、警示框。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT


def _set_cn(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)


def new_doc(title, subtitle=""):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = '宋体'
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    sec = doc.sections[0]
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn(t.add_run(title), "黑体", 22, True, (20, 60, 110))
    if subtitle:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cn(s.add_run(subtitle), "楷体", 12, False, (90, 90, 90))
    return doc


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _set_cn(p.add_run(text), "黑体", 15, True, (20, 80, 140))


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    _set_cn(p.add_run(text), "黑体", 12.5, True, (40, 100, 60))


def para(doc, text, size=11, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    _set_cn(p.add_run(text), "宋体", size)
    return p


def bullet(doc, text, size=11, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18 + level * 18)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(2)
    _set_cn(p.add_run("• " + text), "宋体", size)
    return p


def note(doc, text, color=(170, 40, 40)):
    """警示/提示框（用单格表格模拟底色）"""
    tb = doc.add_table(rows=1, cols=1)
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tb.cell(0, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.3
    _set_cn(p.add_run(text), "楷体", 10.5, True, color)
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'FFF4E5')
    cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()
    return tb


def table(doc, header, rows, widths=None):
    tb = doc.add_table(rows=1, cols=len(header))
    tb.style = 'Table Grid'
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tb.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cn(p.add_run(h), "黑体", 10.5, True, (255, 255, 255))
        from docx.oxml import OxmlElement
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '3A6EA5')
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for r in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            _set_cn(p.add_run(str(v)), "宋体", 10)
    doc.add_paragraph()
    return tb


def figure(doc, path, caption, width=Cm(14.5)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=width)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cn(c.add_run(caption), "楷体", 10, False, (80, 80, 80))
    c.paragraph_format.space_after = Pt(8)


def save(doc, path):
    doc.save(path)
    print("saved", path)
